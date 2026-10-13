import docx
from docx.shared import Inches, Pt
import win32com.client
import pymupdf

doc = docx.Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11.0)
sec.top_margin = Inches(1.0)
sec.bottom_margin = Inches(1.0)
sec.left_margin = Inches(1.0)
sec.right_margin = Inches(1.0)

for i in range(120):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(f'Line {i+1:03d}: const sampleVar = await EvidenceModel.findById(req.params.id);')
    r.font.name = 'Consolas'
    r.font.size = Pt(12)

doc.save(r'C:\Users\asara\OneDrive\Documents\cloudevidence\test_12pt.docx')

word = win32com.client.Dispatch('Word.Application')
word.Visible = False
doc_word = word.Documents.Open(r'C:\Users\asara\OneDrive\Documents\cloudevidence\test_12pt.docx')
doc_word.SaveAs(r'C:\Users\asara\OneDrive\Documents\cloudevidence\test_12pt.pdf', FileFormat=17)
doc_word.Close(False)
word.Quit()

doc_pdf = pymupdf.open(r'C:\Users\asara\OneDrive\Documents\cloudevidence\test_12pt.pdf')
print('Total pages for 120 lines at 12pt:', len(doc_pdf))
for idx, page in enumerate(doc_pdf):
    lines = page.get_text().strip().splitlines()
    blocks = page.get_text('blocks')
    last_block = blocks[-1] if blocks else None
    y1 = last_block[3] if last_block else 0
    print(f'Page {idx+1}: {len(lines)} lines, first="{lines[0]}", last="{lines[-1]}", last y1={y1:.1f}')
