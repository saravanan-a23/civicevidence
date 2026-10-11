import docx

doc = docx.Document(r"C:\Users\asara\Downloads\DOC-20261011-WA0000.docx")
found = False
for p in doc.paragraphs:
    if "APPENDIX C" in p.text.upper() or "APPENDIX D" in p.text.upper():
        found = True
    if found:
        if p.text.strip():
            print(repr(p.text.strip()))

for t in doc.tables:
    first_cell = t.cell(0, 0).text.strip()
    if "SCORE" in first_cell.upper() or "EVALUATION" in first_cell.upper() or "RUBRIC" in first_cell.upper():
        print("Found table in docx:", first_cell)
